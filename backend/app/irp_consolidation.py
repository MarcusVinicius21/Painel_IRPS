"""Versioned internal reports. Only current, documented confirmations enter totals."""
import hashlib,io
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from . import director,irp_register
from .db import identifier
from .exports import safe_text


def export(ctx,runtime,sid,fmt):
    ctx.require('export');ctx.require('document.read')
    study=director.study(ctx,sid);register=irp_register.summary(ctx,sid);p=study['payload']
    cities={r['ibge']:r['name']+'/'+r['uf'] for r in ctx.db.all('SELECT ibge,name,uf FROM municipality')}
    rows=[]
    for r in register['municipalities']:
        proof=r['manifestation']['payload'].get('document_id') if r['manifestation'] else None
        rows.append([cities.get(r['ibge'],r['ibge']),r['ibge'],p['name'],p['unit'],r['suggested_quantity'],
                     r['confirmed_quantity'],r['display_state'],str(proof or ''),
                     r['manifestation']['revision'] if r['manifestation'] else None])
    labels={'communication_pending':'Comunicação pendente','awaiting_response':'Aguardando manifestação',
        'overdue':'Sem manifestação no prazo','received':'Resposta recebida','in_review':'Em conferência',
        'confirmed':'Confirmada','adjusted_confirmed':'Ajustada e confirmada','not_participating':'Não participa',
        'invalid':'Manifestação inválida','requires_reconfirmation':'Exige nova conferência'}
    for r in rows:r[6]=labels.get(r[6],r[6])
    title='Consolidação interna · '+('Demonstração fictícia' if study['scenario']['classification']=='synthetic' else 'Dados internos')
    notice='Somente confirmações documentadas e atuais. Pendências não equivalem a zero. Não aprova, publica ou assina atos.'
    if fmt=='xlsx':
        wb=Workbook();wb.remove(wb.active)
        sheets={'Por município':[['Município','IBGE','Item','Unidade','Sugestão conferida','Quantidade confirmada','Situação','Documento','Versão da manifestação']]+rows,
            'Por item':[['Item','Unidade','Confirmado conhecido','Total consolidado','Municípios confirmados','Pendentes','Não participantes'],
                [p['name'],p['unit'],register['confirmed_known_total'],register['consolidated_total'],register['confirmed_municipalities'],register['pending_municipalities'],register['non_participating_municipalities']]],
            'Pendências':[['Município','IBGE','Situação']]+[[r[0],r[1],r[6]] for r,c in zip(rows,register['municipalities']) if not c['confirmation_current'] and not c['non_participation_current']],
            'Contexto':[['Campo','Valor'],['Relatório',title],['Estudo',str(sid)],['Cenário',study['scenario_id']],
                ['Versão do cenário',study['scenario']['current_revision']],['Contexto atual',str(register['context_current'])],['Regra',notice]]}
        for name,values in sheets.items():
            ws=wb.create_sheet(name)
            for row in values:ws.append([safe_text(v) if isinstance(v,str) else v for v in row])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='005ED9')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(60,max(18,max(len(str(c.value or '')) for c in col)+2))
                for cell in col:cell.alignment=Alignment(vertical='top',wrap_text=True)
        out=io.BytesIO();wb.save(out);data=out.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from .irp_documents import pdf_bytes
        parts=[('title',title),('p',notice),('p',f'Estudo {sid} · cenário versão {study["scenario"]["current_revision"]} · estudo na versão atual: {"sim" if register["context_current"] else "não"}'),
            ('p',f'Item: {p["name"]} · unidade: {p["unit"]}. Confirmado conhecido: {register["confirmed_known_total"] if register["confirmed_known_total"] is not None else "Não disponível"}.'),
            ('table',{'headers':['Município','IBGE','Unidade','Sugerida','Confirmada','Situação'],
                'rows':[[r[0],r[1],r[3],r[4] if r[4] is not None else 'Não disponível',r[5] if r[5] is not None else 'Não disponível',r[6]] for r in rows]})]
        data=pdf_bytes(parts);mime='application/pdf'
    aid=identifier();key=f'exports/{ctx.org_id}/{aid}.{fmt}';path=Path(runtime)/key;path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data)
    filename=f'IRPS_consolidacao_v{study["scenario"]["current_revision"]}.{fmt}';sha=hashlib.sha256(data).hexdigest()
    ctx.db.execute('''INSERT INTO artifact(org_id,id,classification,storage_key,filename,media_type,sha256,scenario_id,revision,created_by)
        VALUES (:o,:i,:c,:s,:f,:m,:h,:sid,:v,:a)''',{'o':ctx.org_id,'i':aid,'c':study['scenario']['classification'],'s':key,'f':filename,'m':mime,'h':sha,'sid':study['scenario_id'],'v':study['scenario']['current_revision'],'a':ctx.actor_id})
    ctx.audit('irp.consolidation_exported',sid,{'artifact_id':aid,'sha256':sha,'context_current':register['context_current']})
    return {'id':aid,'download_url':'/api/v1/artifacts/'+aid+'/download','filename':filename,'sha256':sha,'context_current':register['context_current']}
