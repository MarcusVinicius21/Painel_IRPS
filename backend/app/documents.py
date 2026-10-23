"""Local bounded readers. Content is untrusted evidence, never executable instructions."""
from pathlib import Path,PurePosixPath
from decimal import Decimal,InvalidOperation
import hashlib,io,zipfile,re,csv,unicodedata
from defusedxml import ElementTree as ET
from .security import fail
from .db import identifier,encoded

LIMIT=2*1024*1024
VERSION='irps-local-readers-v2.1'

def normalized(t):return ''.join(x for x in unicodedata.normalize('NFD',t.lower()) if unicodedata.category(x)!='Mn')

def number(raw):
 try:
  d=Decimal(raw.replace('.','').replace(',','.') if ',' in raw else raw)
  if not d.is_finite() or d<0 or d>Decimal('1000000000000'):raise ValueError()
  return str(d)
 except (InvalidOperation,ValueError):return None

def checked_zip(data):
 try:z=zipfile.ZipFile(io.BytesIO(data))
 except zipfile.BadZipFile:fail(422,'Não conseguimos abrir este arquivo. Confira o formato.','BAD_DOCUMENT')
 infos=z.infolist()
 if len(infos)>1000 or sum(i.file_size for i in infos)>20*1024*1024:fail(413,'Divida o arquivo em partes menores.','DOCUMENT_LIMIT')
 for i in infos:
  p=PurePosixPath(i.filename.replace('\\','/'))
  if p.is_absolute() or '..' in p.parts or ':' in p.parts[0] or i.file_size/max(i.compress_size,1)>200:fail(422,'O arquivo precisa de verificação antes da leitura.','UNSAFE_DOCUMENT')
  if any(x in i.filename.lower() for x in ('vbaproject','activex','embeddings/')):fail(422,'Use uma cópia sem macros ou objetos incorporados.','ACTIVE_CONTENT')
  if i.filename.endswith('.rels') and re.search(br'TargetMode\s*=\s*[\"\x27]External',z.read(i),re.I):fail(422,'Use uma cópia sem vínculos externos.','EXTERNAL_REFERENCE')
 return z

def parse(data,extension):
 if not data or len(data)>LIMIT:fail(413,'Escolha um arquivo de até 2 MB.','FILE_SIZE')
 segments=[];fields=[];notices=[]
 if extension=='.pdf':
  if not data.startswith(b'%PDF-'):fail(422,'O conteúdo não corresponde a um PDF.','MIME_MISMATCH')
  from pypdf import PdfReader
  try:
   reader=PdfReader(io.BytesIO(data))
   if reader.is_encrypted:fail(422,'Use uma cópia sem senha.','ENCRYPTED_DOCUMENT')
   if len(reader.pages)>50:fail(413,'Divida o PDF em partes de até 50 páginas.','PAGE_LIMIT')
   for i,p in enumerate(reader.pages):segments.append({'locator':{'page':i+1},'text':(p.extract_text() or '').strip()[:30000]})
   if not any(s['text'] for s in segments):notices.append('Não encontramos texto selecionável. A leitura por imagem ainda está desativada.')
  except __import__('fastapi').HTTPException:raise
  except Exception:fail(422,'Não conseguimos ler este PDF. Confira o original.','PDF_READING')
 elif extension=='.docx':
  z=checked_zip(data)
  if 'word/document.xml' not in z.namelist():fail(422,'O conteúdo não corresponde a um documento DOCX.','MIME_MISMATCH')
  try:tree=ET.fromstring(z.read('word/document.xml'))
  except Exception:fail(422,'Não conseguimos ler a estrutura deste documento.','DOCX_READING')
  ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  for i,p in enumerate(tree.findall('.//w:p',ns)):
   text=''.join(t.text or '' for t in p.findall('.//w:t',ns))
   if text.strip():segments.append({'locator':{'paragraph':i+1},'text':text[:10000]})
 elif extension=='.xlsx':
  z=checked_zip(data)
  if 'xl/workbook.xml' not in z.namelist():fail(422,'O conteúdo não corresponde a uma planilha XLSX.','MIME_MISMATCH')
  from openpyxl import load_workbook
  try:wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
  except Exception:fail(422,'Não conseguimos ler a planilha. Confira o original.','XLSX_READING')
  count=0
  try:
   for sheet in wb:
    if sheet.max_row and sheet.max_column and sheet.max_row*sheet.max_column>200000:fail(413,'Divida a planilha em partes menores.','CELL_LIMIT')
    headers={}
    for row in sheet.iter_rows():
     found={cell.column:normalized(str(cell.value)).strip() for cell in row if cell.value is not None and normalized(str(cell.value)).strip() in ('descricao','unidade','quantidade','preco','valor')}
     if len(found)>=2:headers=found
     for cell in row:
      if cell.value is None:continue
      count+=1
      if count>20000:fail(413,'A planilha excede 20 mil células preenchidas.','CELL_LIMIT')
      loc={'sheet':sheet.title,'cell':cell.coordinate};formula=cell.data_type=='f';raw=str(cell.value)[:10000]
      segments.append({'locator':loc,'text':raw,'formula':formula})
      if formula:fields.append({'name':'formula','value':None,'raw':raw[:1000],'locator':loc,'state':'review_required','reason':'Fórmulas não são executadas. Confira o valor na planilha original.'})
      elif headers and not found and cell.column in headers:
       name=headers[cell.column];val=raw
       if name in ('quantidade','preco','valor'):val=None if isinstance(cell.value,str) and re.fullmatch(r'\d+\.\d{3}',raw) else number(raw)
       fields.append({'name':name,'value':val,'raw':raw,'locator':loc,'state':'review_required','reason':'Confira a coluna, o número e a unidade antes de usar.'})
  finally:wb.close()
 elif extension=='.csv':
  try:text=data.decode('utf-8-sig')
  except UnicodeDecodeError:fail(422,'Salve a planilha no formato CSV UTF-8.','CSV_ENCODING')
  if '\x00' in text:fail(422,'O conteúdo não corresponde a um CSV.','MIME_MISMATCH')
  try:delimiter=csv.Sniffer().sniff(text[:4096],delimiters=',;\t').delimiter
  except csv.Error:delimiter=';'
  count=0;headers={}
  for r,row in enumerate(csv.reader(io.StringIO(text),delimiter=delimiter)):
   found={c:normalized(raw).strip() for c,raw in enumerate(row) if normalized(raw).strip() in ('descricao','unidade','quantidade','preco','valor')}
   if len(found)>=2:headers=found
   for c,raw in enumerate(row):
    count+=1
    if count>20000:fail(413,'Divida a planilha em partes menores.','CELL_LIMIT')
    formula=raw.lstrip().startswith(('=','+','-','@'));loc={'row':r+1,'column':c+1}
    segments.append({'locator':loc,'text':raw[:10000],'formula':formula})
    if formula:fields.append({'name':'formula','value':None,'raw':raw[:1000],'locator':loc,'state':'review_required','reason':'Conteúdo interpretável como fórmula. Confira o original.'})
    elif headers and not found and c in headers:
     name=headers[c];value=raw
     if name in ('quantidade','preco','valor'):value=None if re.fullmatch(r'\d+\.\d{3}',raw) else number(raw)
     fields.append({'name':name,'value':value,'raw':raw,'locator':loc,'state':'review_required','reason':'Confira a coluna, o separador decimal e a unidade.'})
 else:fail(422,'Use PDF, DOCX, XLSX ou CSV.','UNSUPPORTED_FORMAT')
 if sum(len(s['text']) for s in segments)>500000:fail(413,'Divida o documento em partes menores.','TEXT_LIMIT')
 unsafe=any(re.search(r'ignore (?:as |todas as )?(?:regras|instru)|envie (?:a |as )?(?:chave|senha)|reveal.*(?:key|secret)',normalized(s['text'])) for s in segments)
 if unsafe:notices.append('Encontramos instruções no documento. Elas não foram executadas.')
 for s in segments:
  if s.get('formula'):continue
  for m in re.finditer(r'(?i)\b(quantidade|preço|preco|valor)\s*:\s*(?:R\$\s*)?([\d.,]+)',s['text']):fields.append({'name':normalized(m.group(1)),'value':number(m.group(2)),'raw':m.group(2),'locator':s['locator'],'state':'review_required'})
 return {'segments':segments,'fields':fields,'notices':notices,'instruction_detected':unsafe,'state':'needs_review','external_ai_used':False,'parser_version':VERSION,'antivirus_scan':'not_available','applied_to_scenario':False}

def ingest(ctx,runtime,filename,data,classification):
 if ctx.operational and classification=='synthetic':fail(422,'O ambiente operacional não aceita documentos fictícios.','SYNTHETIC_DISABLED')
 ctx.require('document.write')
 if classification=='business_private':ctx.require('quote.write')
 parsed=parse(data,Path(filename).suffix.lower());sha=hashlib.sha256(data).hexdigest()
 ctx.db.execute('SELECT pg_advisory_xact_lock(hashtextextended(:k,0))',{'k':'document-ingest:'+ctx.org_id+':'+sha})
 existing=ctx.db.one('SELECT id FROM document WHERE org_id=:o AND sha256=:s',{'o':ctx.org_id,'s':sha})
 if existing:
  did=str(existing['id']);ctx.db.execute('SELECT pg_advisory_xact_lock(hashtextextended(:k,0))',{'k':'document-review:'+ctx.org_id+':'+did});previous=get(ctx,did)
  if previous['parser_version']!=VERSION:
   ctx.db.execute("INSERT INTO extracted_document(org_id,document_id,revision,parser_version,result,review_state) VALUES (:o,:i,:v,:p,CAST(:r AS jsonb),'needs_review')",{'o':ctx.org_id,'i':did,'v':previous['revision']+1,'p':VERSION,'r':encoded(parsed)});ctx.audit('document.reprocessed',did,{'revision':previous['revision']+1})
  return {**get(ctx,did),'duplicate':True}
 did=identifier();key='uploads/'+ctx.org_id+'/'+sha+Path(filename).suffix.lower();path=Path(runtime)/key;path.parent.mkdir(parents=True,exist_ok=True)
 if path.exists():
  if hashlib.sha256(path.read_bytes()).hexdigest()!=sha:fail(409,'O original precisa de verificação.','ORIGINAL_CONFLICT')
 else:
  with path.open('xb') as f:f.write(data)
 ctx.db.execute('INSERT INTO document(org_id,id,filename,sha256,storage_key,classification,created_by) VALUES (:o,:i,:f,:h,:s,:c,:a)',{'o':ctx.org_id,'i':did,'f':Path(filename.replace('\\','/')).name[:180],'h':sha,'s':key,'c':classification,'a':ctx.actor_id})
 ctx.db.execute("INSERT INTO extracted_document(org_id,document_id,revision,parser_version,result,review_state) VALUES (:o,:i,1,:p,CAST(:r AS jsonb),'needs_review')",{'o':ctx.org_id,'i':did,'p':VERSION,'r':encoded(parsed)})
 ctx.audit('document.extracted',did,{'bytes':len(data),'fields':len(parsed['fields'])});return {**get(ctx,did),'duplicate':False}

def get(ctx,did):
 ctx.require('document.read');row=ctx.db.one('SELECT d.*,e.revision,e.result,e.review_state,e.parser_version FROM document d JOIN extracted_document e ON e.org_id=d.org_id AND e.document_id=d.id WHERE d.org_id=:o AND d.id=:i ORDER BY e.revision DESC LIMIT 1',{'o':ctx.org_id,'i':did})
 if not row:fail(404,'Este documento não está disponível neste espaço.','NOT_FOUND')
 if row['classification']=='business_private':ctx.require('quote.read')
 return row

def review(ctx,did,body):
 ctx.require('document.review');ctx.db.execute('SELECT pg_advisory_xact_lock(hashtextextended(:k,0))',{'k':'document-review:'+ctx.org_id+':'+did});row=get(ctx,did)
 if row['revision']!=body.revision:fail(409,'O documento mudou. Abra a revisão atual.','VERSION_CONFLICT')
 if row['result']['instruction_detected'] and body.decision=='reviewed':fail(422,'Revise o conteúdo suspeito antes de concluir.','SUSPICIOUS_CONTENT')
 result={**row['result'],'review_note':body.note,'proposed_corrections':body.corrections,'applied_to_scenario':False}
 ctx.db.execute('INSERT INTO extracted_document(org_id,document_id,revision,parser_version,result,review_state,reviewer) VALUES (:o,:i,:v,:p,CAST(:r AS jsonb),:s,:a)',{'o':ctx.org_id,'i':did,'v':row['revision']+1,'p':VERSION,'r':encoded(result),'s':body.decision,'a':ctx.actor_id});ctx.audit('document.reviewed',did,{'revision':row['revision']+1});return get(ctx,did)
