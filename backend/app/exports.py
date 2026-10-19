from pathlib import Path
from decimal import Decimal
import csv,io,hashlib
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape,A4
from .db import identifier,encoded
from .security import fail
from .scenarios import get_scenario

def safe_text(s):
 s=str(s or '')
 return "'"+s if s.lstrip().startswith(('=','+','-','@','\t','\r')) else s

def create(ctx,runtime,sid,fmt,calculation_id):
 ctx.require('export');scenario=get_scenario(ctx,sid)
 calc=ctx.db.one('SELECT * FROM calculation WHERE org_id=:o AND id=:i AND scenario_id=:s',{'o':ctx.org_id,'i':calculation_id,'s':sid})
 if not calc:fail(404,'Este cálculo não está disponível no cenário.','NOT_FOUND')
 rows=ctx.db.all('SELECT l.*,m.name AS municipality_name,i.name AS item_name FROM calculation_line l JOIN municipality m ON m.ibge=l.ibge JOIN item_version i ON i.id=l.item_id AND i.version=l.item_version WHERE l.org_id=:o AND calculation_id=:c ORDER BY l.ibge,l.item_id',{'o':ctx.org_id,'c':calculation_id})
 headers=['Município','Código IBGE','Item','Versão','Unidade','Quantidade','Preço de referência (R$)','Valor estimado (R$)','Situação','Memória','Origem do preço']
 records=[[r['municipality_name'],r['ibge'],r['item_name'],str(r['item_version']),r['unit'],str(r['quantity']) if r['quantity'] is not None else 'Não disponível',str(r['unit_price']) if r['unit_price'] is not None else 'Não disponível',str(r['value']) if r['value'] is not None else 'Não disponível',r['state'],r['memory']['formula'],(r['memory'].get('price_source') or {}).get('source_url','Não disponível')] for r in rows]
 title=('DADOS FICTÍCIOS · ' if scenario['classification']=='synthetic' else 'PLANEJAMENTO INTERNO · ')+scenario['title']
 subtitle=f"Versão {calc['revision']} | {calc['summary']['period']} | Quantidades para revisão. Não autoriza despesa."
 if fmt=='csv':
  output=io.StringIO(newline='');writer=csv.writer(output,delimiter=';');writer.writerow(['Classificação',title]);writer.writerow(['Escopo',subtitle]);writer.writerow(headers)
  writer.writerows([[safe_text(c) for c in row] for row in records]);data=b'\xef\xbb\xbf'+output.getvalue().encode('utf-8');mime='text/csv; charset=utf-8'
 elif fmt=='xlsx':
  wb=Workbook();ws=wb.active;ws.title='Quantidades e valores';ws.append([title]);ws.append([subtitle]);ws.append(headers)
  for row in records:
   converted=[safe_text(c) for c in row]
   ws.append(converted)
  # Exact decimal values remain text, preventing spreadsheet float round-trip loss.
  ws.freeze_panes='F4';ws.auto_filter.ref=f'A3:K{len(records)+3}'
  for col in 'ABCDEFGHIJK':ws.column_dimensions[col].width=24 if col not in ('C','J','K') else 50
  from openpyxl.styles import Font,PatternFill
  for cell in ws[3]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='006772')
  meta=wb.create_sheet('Metodologia');meta.append(['Núcleo V4 SHA-256',calc['core_sha256']]);meta.append(['Valor total estimado',calc['summary']['total'] or 'Pendente']);meta.append(['Dados fictícios',scenario['classification']=='synthetic']);meta.append(['Aviso','Nenhum valor deste cenário comprova consumo ou pagamento.'])
  output=io.BytesIO();wb.save(output);data=output.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 elif fmt=='pdf':
  output=io.BytesIO();pdf=canvas.Canvas(output,pagesize=landscape(A4));width,height=landscape(A4);page=0
  def heading():
   nonlocal page
   page+=1;pdf.setFillColorRGB(.02,.25,.29);pdf.setFont('Helvetica-Bold',13);pdf.drawString(30,height-32,'IRPS | CONSORMULTI');pdf.setFont('Helvetica',9);pdf.drawString(30,height-49,title[:140]);pdf.drawString(30,height-64,subtitle[:160]);pdf.drawString(30,20,f'Página {page} | Dados da versão preservada | Cálculos V4')
  heading();y=height-90
  for r,record in zip(rows,records):
   if y<60:pdf.showPage();heading();y=height-90
   pdf.setFont('Helvetica-Bold',8);pdf.drawString(30,y,(record[0]+' · '+record[1]+' | '+record[2])[:145]);y-=13;pdf.setFont('Helvetica',8)
   pdf.drawString(30,y,f"v{record[3]} | {record[4]} | Quantidade: {record[5]} | Preço: R$ {record[6]} | Estimado: R$ {record[7]}");y-=12
   pdf.drawString(30,y,('Cálculo: '+record[9])[:170]);y-=12
   pdf.drawString(30,y,('Fonte: '+record[10])[:170]);y-=19
  pdf.save();data=output.getvalue();mime='application/pdf'
 else:fail(422,'Escolha CSV, XLSX ou PDF.')
 aid=identifier();filename=f'IRPS_cenario_v{calc["revision"]}.{fmt}';key=f'exports/{ctx.org_id}/{aid}.{fmt}';path=Path(runtime)/key;path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data);sha=hashlib.sha256(data).hexdigest()
 ctx.db.execute('INSERT INTO artifact(org_id,id,classification,storage_key,filename,media_type,sha256,scenario_id,revision,created_by) VALUES (:o,:i,:c,:s,:f,:m,:h,:sid,:v,:a)',{'o':ctx.org_id,'i':aid,'c':scenario['classification'],'s':key,'f':filename,'m':mime,'h':sha,'sid':sid,'v':calc['revision'],'a':ctx.actor_id});ctx.audit('scenario.exported',sid,{'artifact_id':aid,'format':fmt,'lines':len(rows),'revision':calc['revision']})
 return {'id':aid,'filename':filename,'rows':len(rows),'bytes':len(data),'sha256':sha,'revision':calc['revision'],'download_url':'/api/v1/artifacts/'+aid+'/download'}

def authorized_file(ctx,runtime,aid):
 ctx.require('export');r=ctx.db.one('SELECT * FROM artifact WHERE org_id=:o AND id=:i',{'o':ctx.org_id,'i':aid})
 if not r:fail(404,'Este arquivo não está disponível neste espaço.','NOT_FOUND')
 if r['classification']=='business_private':ctx.require('quote.read')
 root=Path(runtime).resolve();path=(root/r['storage_key']).resolve()
 if root not in path.parents or not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest()!=r['sha256']:fail(409,'O arquivo precisa ser gerado novamente.','ARTIFACT_CONFLICT')
 return path,r
