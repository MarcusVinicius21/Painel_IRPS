"""Approved journey: truthful executive data and evidence-bound internal consolidation."""
from datetime import date,timedelta
from io import BytesIO
from openpyxl import load_workbook
from pypdf import PdfReader
from test_director_workflow import make_study,draft_body,national
from test_irp_register import doc,review,operation


def test_dashboard_independent_sections_are_official_and_organization_scoped(login):
    c=login();s=make_study(c)
    cat=c.get('/api/v1/director/dashboard/catalogue');assert cat.status_code==200,cat.text
    assert cat.json()['purchase_history_implied'] is False
    h=c.get('/api/v1/director/dashboard/history');assert h.status_code==200,h.text
    assert all(r['provider']!='DEMO' for r in h.json()['phases'])
    assert all(r['month'] for r in h.json()['monthly'])
    a=c.get('/api/v1/director/dashboard/operations');assert a.status_code==200,a.text
    assert s['id'] in [r['id'] for r in a.json()['studies']]
    other=login('compras-b').get('/api/v1/director/dashboard/operations').json()
    assert s['id'] not in [r['id'] for r in other['studies']]
    assert c.get('/api/v1/director/dashboard/unknown').status_code==422


def test_saved_drafts_and_exact_catalogue_selection_can_be_reopened(login,national):
    c=login();snap,_,prefix=national
    entry=c.get('/api/v1/national-catalogue/'+snap['snapshot_id']+'/'+prefix+'0');assert entry.status_code==200
    assert entry.json()['code']==prefix+'0'
    s=make_study(c);url='/api/v1/director/studies/'+s['id']+'/drafts'
    d=c.post(url,json=draft_body());assert d.status_code==201
    assert c.get(url).json()['records'][0]['id']==d.json()['id']
    assert login('compras-b').get(url).status_code==404
    preview=c.get('/api/v1/irp-drafts/'+d.json()['id']+'/preview')
    assert preview.status_code==200 and preview.json()['sections']
    assert not preview.json()['stale']
    assert login('compras-b').get('/api/v1/irp-drafts/'+d.json()['id']+'/preview').status_code==404


def test_pending_failed_document_and_refusal_have_distinct_states(login):
    c=login();s=make_study(c);url='/api/v1/director/studies/'+s['id']
    rows=c.get(url+'/register').json()['municipalities'];assert all(r['display_state']=='communication_pending' for r in rows)
    r=c.post(url+'/operations',json=operation(kind='communication',status='prepared',quantity=None,deadline=(date.today()-timedelta(days=1)).isoformat()))
    assert r.status_code==201,r.text
    rows=c.get(url+'/register').json()['municipalities'];assert next(r for r in rows if r['ibge']=='3131307')['display_state']=='overdue'
    d=doc(c);review(c,d['id'],d['revision'])
    r=c.post(url+'/operations',json=operation(status='confirmed',quantity='30',document_id=d['id']));assert r.status_code==201,r.text
    r=c.post(url+'/operations',json=operation(municipality_ibge='3101607',status='not_participating',quantity=None,document_id=d['id']));assert r.status_code==201,r.text
    reg=c.get(url+'/register').json()
    assert reg['consolidated_total']=='30' and reg['confirmed_municipalities']==1 and reg['non_participating_municipalities']==1
    assert reg['all_invitations_resolved'] and not reg['all_municipalities_confirmed'] and reg['pending_municipalities']==0
    assert len(reg['history'])==3
    # A review change invalidates confirmation AND refusal, rather than counting them as zero.
    current=c.get('/api/v1/documents').json()['records'];document=next(r for r in current if r['id']==d['id'])
    review(c,d['id'],document['revision'],'needs_changes')
    reg=c.get(url+'/register').json();assert reg['consolidated_total'] is None and reg['non_participating_municipalities']==0
    assert all(r['display_state']=='requires_reconfirmation' for r in reg['municipalities'])


def test_consolidation_exports_keep_pending_null_and_exclude_private_data(login):
    c=login();s=make_study(c,name='=HYPERLINK item ficticio de teste');url='/api/v1/director/studies/'+s['id']
    d=doc(c);review(c,d['id'],d['revision'])
    assert c.post(url+'/operations',json=operation(document_id=d['id'],quantity='0')).status_code==201
    r=c.post(url+'/consolidation/xlsx');assert r.status_code==201,r.text
    raw=c.get(r.json()['download_url']);assert raw.status_code==200
    wb=load_workbook(BytesIO(raw.content));assert wb.sheetnames==['Por município','Por item','Pendências','Contexto']
    rows=list(wb['Por município'].values)
    assert all(row[2].startswith("'=") for row in rows[1:])
    assert sorted([row[5] for row in rows[1:] if row[5] is not None])==['0']
    assert any(row[5] is None for row in rows[1:])
    assert 'business_private' not in str(rows) and 'password' not in str(rows)
    p=c.post(url+'/consolidation/pdf');assert p.status_code==201,p.text
    pdf=PdfReader(BytesIO(c.get(p.json()['download_url']).content));assert len(pdf.pages)>=1
    text=''.join(page.extract_text() for page in pdf.pages)
    assert 'Somente' in text and 'True' not in text and 'False' not in text
    other=login('compras-b');assert other.post(url+'/consolidation/xlsx').status_code==404
    assert other.get(r.json()['download_url']).status_code==404
    # Existing read-only role explicitly has export permission; preserve that contract.
    viewer=login('consulta-a')
    assert viewer.post(url+'/consolidation/xlsx').status_code==201
    assert viewer.post(url+'/operations',json=operation()).status_code==403


def test_incomplete_progress_is_immutable_isolated_and_never_creates_calculations(login,national):
    c=login();snap,_,prefix=national
    before=c.get('/api/v1/scenarios').json()
    body={'snapshot_id':snap['snapshot_id'],'code':prefix+'0','municipalities':['3131307'],
        'start':'2025-01-01','end':'2025-12-31','phase':'homologado','stage':'planejamento',
        'form':{'name':'Objeto em revisão','category':'veiculo','unit':'','preliminary_target':'200'}}
    p=c.post('/api/v1/director/progress',json=body);assert p.status_code==201,p.text
    saved=p.json();payload=saved['payload']
    assert saved['draft_only'] and not saved['approval_implied']
    assert payload['form']['unit']=='' and not payload['technical_confirmed'] and not payload['v4_input_created']
    assert c.get('/api/v1/scenarios').json()==before
    assert c.get('/api/v1/director/progress/'+saved['id']).json()['payload']==payload
    assert login('compras-b').get('/api/v1/director/progress/'+saved['id']).status_code==404
    assert login('consulta-a').post('/api/v1/director/progress',json=body).status_code==403
    assert c.post('/api/v1/director/progress',json={**body,'technical_confirmed':True}).status_code==422
    revision=c.post('/api/v1/director/progress',json={**body,'supersedes_id':saved['id'],'form':{**body['form'],'unit':'unidade'}})
    assert revision.status_code==201,revision.text
    assert c.post('/api/v1/director/progress',json={**body,'supersedes_id':saved['id']}).status_code==409
    assert c.get('/api/v1/director/progress/'+saved['id']).json()['payload']==payload
    assert saved['id'] not in [r['id'] for r in c.get('/api/v1/director/progress').json()['records']]
    assert revision.json()['id'] in [r['id'] for r in c.get('/api/v1/director/progress').json()['records']]
