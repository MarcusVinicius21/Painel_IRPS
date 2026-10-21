"""Real PostgreSQL tests for national snapshots, private studies and safe drafts."""
import csv,io,zipfile
from uuid import uuid4
from datetime import date
import pytest
from app.national_catalogue import import_csv
from app import catalogue,irp_documents


@pytest.fixture
def national(settings,tmp_path):
    path=tmp_path/'materials.csv'
    prefix=str(int(uuid4().hex[:8],16))
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f,delimiter=';');w.writerow(['codigoGrupo','nomeGrupo','codigoClasse','nomeClasse','codigoPdm','nomePdm','codigoItem','descricaoItem','statusItem'])
        for n in range(56):w.writerow(['89','Alimentos','8950','Café','1','Café',prefix+str(n),f'Café torrado e moído tipo {n}' if n%2==0 else f'CAFÉ TORRADO E MOÍDO TIPO {n}',str(n%2==0)])
    result=import_csv(settings.owner_database_url.get_secret_value(),path,'CATMAT')
    return result,path,prefix


def study_body():
    return {'classification':'synthetic','name':'Veículo de passeio para estudo local','category':'veiculo','specification':'Veículo de passeio cinco lugares, especificação de teste para revisão.',
        'unit':'unidade','quantity_step':'1','municipalities':['3131307','3101607'],'phase':'homologado',
        'start':'2025-01-01','end':'2025-12-31','preliminary_target':'200',
        'target_justification':'Meta exploratória informada explicitamente para o estudo.',
        'technical_evidence':'Referência técnica fictícia do teste local.','technical_confirmed':True}


def make_study(client,**updates):
    r=client.post('/api/v1/director/studies',json={**study_body(),**updates})
    assert r.status_code==201,r.text
    return r.json()


def draft_body(**updates):
    return {'scenario_revision':1,'object':'Aquisição eventual de veículo de passeio para estudo.','coherence_confirmed':True,
        'coherence_evidence':'Conferência fictícia de coerência no teste local.',**updates}


def test_snapshot_import_idempotent_atomic_and_paged(login,settings,national,tmp_path):
    result,path,prefix=national
    assert result['rows']==56 and result['active']==28 and result['inactive']==28
    assert import_csv(settings.owner_database_url.get_secret_value(),path,'CATMAT')['duplicate']
    c=login();base='/api/v1/national-catalogue';args={'snapshot':result['snapshot_id'],'limit':25}
    r=c.get(base,params=args);assert r.status_code==200,r.text
    first=r.json();assert len(first['records'])==25 and first['next_offset']==25
    second=c.get(base,params={**args,'offset':25}).json();assert not {r['code'] for r in first['records']} & {r['code'] for r in second['records']}
    bycode=c.get(base,params={**args,'q':prefix+'0'}).json()['records'];assert len(bycode)==1 and bycode[0]['name']=='Café torrado e moído tipo 0'
    active=c.get(base,params={**args,'active':'false','limit':50}).json();assert len(active['records'])==28
    assert c.get(base,params={**args,'q':'café torrado','group':'89','classification':'8950'}).json()['records']
    accented=c.get(base,params={**args,'q':'café torrado','limit':50}).json()['records']
    upper=c.get(base,params={**args,'q':'CAFE TORRADO','limit':50}).json()['records']
    assert {r['code'] for r in accented}=={r['code'] for r in upper} and len(accented)==50
    assert any('CAFÉ TORRADO' in r['name'] for r in accented)
    path.write_text('codigoItem;descricaoItem;statusItem\n1;Café;True\n1;Duplicado;True\n',encoding='utf-8')
    with pytest.raises(Exception):import_csv(settings.owner_database_url.get_secret_value(),path,'CATMAT')
    assert len(c.get(base,params={**args,'q':prefix+'0'}).json()['records'])==1
    path.write_text('codigoItem;descricaoItem;statusItem\n1;Café;MAYBE\n',encoding='utf-8')
    with pytest.raises(ValueError):import_csv(settings.owner_database_url.get_secret_value(),path,'CATMAT')


def test_service_csv_preserves_codes_status_and_original(login,settings,tmp_path):
    path=tmp_path/'services.csv';code='000'+str(int(uuid4().hex[:8],16))
    path.write_text('codigoGrupo;nomeGrupo;codigoClasse;nomeClasse;codigoServico;nomeServico;statusServico\n1;Serviços;2;Tecnologia;'+code+';Licença SaaS;False\n',encoding='utf-8')
    r=import_csv(settings.owner_database_url.get_secret_value(),path,'CATSER');c=login()
    rows=c.get('/api/v1/national-catalogue',params={'snapshot':r['snapshot_id'],'q':code}).json()['records']
    assert len(rows)==1 and rows[0]['code']==code and not rows[0]['active']
    assert c.get('/api/v1/national-catalogue/'+r['snapshot_id']+'/'+code+'/links').json()['records']==[]


def test_target_is_not_need_and_private_study_is_isolated(login,national):
    c=login();result,path,prefix=national
    s=make_study(c,snapshot_id=result['snapshot_id'],code=prefix+'0')
    assert s['payload']['preliminary_target']=='200' and s['payload']['history'] is None
    scenario=s['scenario'];assert scenario['snapshot']['inputs']=={} and scenario['snapshot']['allocations']=={}
    calc=c.post('/api/v1/scenarios/'+scenario['id']+'/calculate').json()
    assert calc['summary']['total'] is None and calc['summary']['blocked_lines']==2
    other=login('compras-b')
    assert other.get('/api/v1/director/studies/'+s['id']).status_code==404
    visible=other.get('/api/v1/items').json()
    assert s['payload']['planning_item']['id'] not in {r['id'] for r in visible['versions']}
    assert login('consulta-a').post('/api/v1/director/studies',json=study_body()).status_code==403


@pytest.mark.parametrize('changes',[
    {'unit':'R$/mês'}, {'unit':''}, {'category':'saas','service_period':'','quantity_basis':''},
    {'preliminary_target':'200','target_justification':''}, {'municipalities':['9999999']},
    {'phase':'inventada'}, {'technical_confirmed':False},
])
def test_invalid_definition_cannot_create_study(login,changes):
    r=login().post('/api/v1/director/studies',json={**study_body(),**changes});assert r.status_code==422,r.text


@pytest.mark.parametrize('kind',['procedimento','formulario','planilha','especificacoes','municipios','quantitativos'])
def test_draft_and_exports_are_revision_bound_and_not_official(login,kind):
    c=login();s=make_study(c);r=c.post('/api/v1/director/studies/'+s['id']+'/drafts',json=draft_body(kind=kind,municipality_ibge='3131307'))
    assert r.status_code==201,r.text
    d=r.json();assert d['validation']['state']=='needs_review' and not d['validation']['officially_published']
    assert not d['payload']['signature_invented'] and not d['payload']['legal_clauses_included']
    for fmt in ('docx','pdf'):
        artifact=c.post('/api/v1/irp-drafts/'+d['id']+'/exports/'+fmt)
        assert artifact.status_code==201,artifact.text
        downloaded=c.get(artifact.json()['download_url']);assert downloaded.status_code==200
        if fmt=='docx':
            with zipfile.ZipFile(io.BytesIO(downloaded.content)) as z:
                xml=z.read('word/document.xml').decode();assert 'MINUTA EM REVISÃO' in xml and 'Não confirmada' in xml
                assert 'Conferência fictícia' in xml
        else:
            from pypdf import PdfReader
            text=' '.join(('\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(downloaded.content)).pages)).split())
            assert 'MINUTA EM REVISÃO' in text and 'Não confirmada' in text
    assert login('compras-b').get('/api/v1/irp-drafts/'+d['id']).status_code==404
    old=c.get('/api/v1/scenarios/'+s['scenario']['id']).json()
    updated=c.put('/api/v1/scenarios/'+old['id'],headers={'If-Match':'1'},json={**old['snapshot'],'period':'2027 alterado'})
    assert updated.status_code==200
    assert c.post('/api/v1/irp-drafts/'+d['id']+'/exports/docx').status_code==409


def test_contaminated_object_and_missing_review_block_export(login):
    c=login();s=make_study(c)
    for changes in ({'object':'Locação de plataforma com reconhecimento facial e catracas.'},{'coherence_confirmed':False},{'kind':'formulario','municipality_ibge':'3106200'}):
        r=c.post('/api/v1/director/studies/'+s['id']+'/drafts',json=draft_body(**changes));assert r.status_code==201,r.text
        assert r.json()['validation']['state']=='blocked'
        assert c.post('/api/v1/irp-drafts/'+r.json()['id']+'/exports/docx').status_code==422


@pytest.mark.parametrize('category,unit,description',[
    ('material','kg','Café torrado e moído embalagem de um quilograma.'),
    ('uniforme','unidade','Camiseta escolar algodão, especificação sintética de teste.'),
    ('saas','licença-mês','Licença de software como serviço, especificação de teste.'),
    ('servico','homem-hora','Serviço continuado de suporte técnico, especificação de teste.')])
def test_different_categories_do_not_inherit_vehicle_text(login,category,unit,description):
    c=login();s=make_study(c,name=description,specification=description,category=category,unit=unit,service_period='12 meses' if category in ('saas','servico') else '',quantity_basis='Uma licença ou hora de serviço conforme especificação' if category in ('saas','servico') else '')
    r=c.post('/api/v1/director/studies/'+s['id']+'/drafts',json=draft_body(object=description));assert r.status_code==201
    assert r.json()['validation']['state']=='needs_review'
    assert c.post('/api/v1/irp-drafts/'+r.json()['id']+'/exports/docx').status_code==201


def test_derived_template_uses_requested_target_and_excludes_original_content(tmp_path):
    from docx import Document
    original=tmp_path/'original.docx';target=tmp_path/'derived.docx'
    d=Document();d.add_paragraph('CORPO ORIGINAL CONTAMINADO NÃO COPIAR')
    d.sections[0].header.paragraphs[0].text='CONSORMULTI referência sintética'
    d.sections[0].footer.paragraphs[0].text='CONTATO NÃO VALIDADO'
    d.save(original);before=original.read_bytes()
    result=irp_documents.sanitize_letterhead(original,target)
    assert target.is_file() and original.read_bytes()==before and result['original_body_excluded']
    with zipfile.ZipFile(target) as z:
        assert 'CONTAMINADO' not in z.read('word/document.xml').decode()
        assert all('CONTATO NÃO VALIDADO' not in z.read(name).decode() for name in z.namelist() if name.startswith('word/footer') and name.endswith('.xml'))
    data=irp_documents.docx_bytes([('title','Minuta de teste'),('table',{'headers':['A','B','C','D','E','F'],'rows':[['1','2','3','4','5','6']]})],target)
    document=Document(io.BytesIO(data));assert len(document.tables)==1 and len(document.tables[0].columns)==6
