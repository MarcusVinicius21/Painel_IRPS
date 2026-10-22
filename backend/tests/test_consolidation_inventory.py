import io,json,zipfile
from openpyxl import load_workbook
from pypdf import PdfReader
from test_director_workflow import make_study,draft_body
from test_irp_register import doc,review,operation


def confirm(c,base,code='3131307',quantity='30',status='confirmed'):
 d=doc(c);review(c,d['id'],d['revision']);r=c.post(base+'/operations',json=operation(municipality_ibge=code,document_id=d['id'],quantity=quantity,status=status));assert r.status_code==201,r.text
 return d


def test_partial_and_resolved_totals_and_all_exports_have_traceable_inventory(login):
 c=login();s=make_study(c);base='/api/v1/director/studies/'+s['id'];confirm(c,base)
 r=c.get(base+'/register').json();assert r['confirmed_known_total']=='30' and r['consolidated_total'] is None and r['pending_municipalities']==1
 d=c.post(base+'/drafts',json=draft_body()).json();before=c.get('/api/v1/irp-drafts/'+d['id']).json();inv=c.get(base+'/dossier/inventory');assert inv.status_code==200,inv.text
 inv=inv.json();assert inv['draft_versions_count']==1 and inv['current_drafts'][0]['state']=='draft_needs_review'
 assert inv['current_drafts'][0]['included_if_exported_now'] and inv['official_final_count'] is None
 assert inv['current_drafts'][0]['official_final'] is False and len(inv['supporting_documents'])==1
 assert c.get('/api/v1/irp-drafts/'+d['id']).json()['stale']==before['stale']
 x=c.post(base+'/consolidation/xlsx').json();wb=load_workbook(io.BytesIO(c.get(x['download_url']).content));assert str(wb['Por item'].cell(2,3).value)=='30' and wb['Por item'].cell(2,4).value is None
 pdf=c.post(base+'/consolidation/pdf').json();assert '30' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(c.get(pdf['download_url']).content)).pages)
 package=c.post(base+'/dossier');assert package.status_code==201
 with zipfile.ZipFile(io.BytesIO(c.get(package.json()['download_url']).content)) as z:
  assert json.loads(z.read('indice.json'))==package.json()['manifest']
  assert 'consolidacao/quadro-municipal-e-item.xlsx' in z.namelist()
  assert not any(n.startswith(('uploads/','negocio/')) for n in z.namelist())
 confirm(c,base,'3101607','20');summary=c.get(base+'/register').json();assert summary['consolidated_total']=='50' and summary['all_invitations_resolved']
 inv=c.get(base+'/dossier/inventory').json();assert inv['current_drafts'][0]['state']=='blocked' and not inv['current_drafts'][0]['included_if_exported_now']
 assert inv['generated_artifacts'] and all(a['download_url'].startswith('/api/v1/artifacts/') for a in inv['generated_artifacts'])
 assert len(inv['supporting_documents'])==2
 assert login('compras-b').get(base+'/dossier/inventory').status_code==404
 assert login('consulta-a').get(base+'/dossier/inventory').status_code==200


def test_zero_is_documented_zero_and_decline_never_supplies_quantity(login):
 c=login();s=make_study(c);base='/api/v1/director/studies/'+s['id'];confirm(c,base,quantity='0');confirm(c,base,'3101607',quantity=None,status='not_participating')
 data=c.get(base+'/register').json();assert data['consolidated_total']=='0' and data['confirmed_known_total']=='0' and data['all_invitations_resolved']
 x=c.post(base+'/consolidation/xlsx').json();wb=load_workbook(io.BytesIO(c.get(x['download_url']).content));assert str(wb['Por item'].cell(2,4).value)=='0'
 pdf=c.post(base+'/consolidation/pdf').json();text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(c.get(pdf['download_url']).content)).pages);assert 'Confirmado conhecido: 0' in text


def test_changes_requested_are_excluded_and_inventory_does_not_approve(login):
 c=login();s=make_study(c);base='/api/v1/director/studies/'+s['id'];d=c.post(base+'/drafts',json=draft_body()).json();url='/api/v1/irp-drafts/'+d['id']+'/review'
 def body(m,status):return {'expected_revision':m['revision'],'context_sha256':m['context_sha256'],'status':status,'note':'Parecer sintético para teste de inventário.'}
 m=c.get(url).json();m=c.post(url,json=body(m,'requested')).json();assert c.post(url,json=body(m,'changes_requested')).status_code==201
 inventory=c.get(base+'/dossier/inventory').json();assert inventory['current_drafts'][0]['state']=='blocked'
 package=c.post(base+'/dossier').json();assert package['drafts_included']==0 and any(f.get('included') is False for f in package['manifest']['files'])
 assert c.get(base+'/register').json()['legal_procedure_approved'] is False
